"""Script to generate professional PowerPoint presentation for ACE Mobile Agent System."""

import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_deck():
    prs = Presentation()
    # 16:9 Widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Colors
    NAVY = RGBColor(24, 34, 53)
    ORANGE = RGBColor(255, 107, 53)
    DARK_TEXT = RGBColor(30, 41, 59)
    MUTED_TEXT = RGBColor(100, 116, 139)
    LIGHT_BG = RGBColor(245, 247, 250)
    WHITE = RGBColor(255, 255, 255)
    CARD_BORDER = RGBColor(226, 232, 240)
    ACCENT_BLUE = RGBColor(37, 99, 235)
    GREEN = RGBColor(16, 185, 129)

    def set_bg(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="ACE - INTELLIGENT MOBILE AGENT SYSTEM"):
        # Header banner
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.5), Inches(0.4))
        tf_c = cat_box.text_frame
        tf_c.word_wrap = True
        p_c = tf_c.paragraphs[0]
        p_c.text = category_text.upper()
        p_c.font.size = Pt(11)
        p_c.font.bold = True
        p_c.font.color.rgb = ORANGE

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.5), Inches(0.8))
        tf_t = title_box.text_frame
        tf_t.word_wrap = True
        p_t = tf_t.paragraphs[0]
        p_t.text = title_text
        p_t.font.size = Pt(24)
        p_t.font.bold = True
        p_t.font.color.rgb = NAVY

    def add_card(slide, left, top, width, height, bg_color=WHITE, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    # ==========================================
    # SLIDE 1: Title Slide (Dark Theme)
    # ==========================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1, NAVY)

    # Accent bar
    bar = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.5), Inches(1.2), Inches(0.08))
    bar.fill.solid()
    bar.fill.fore_color.rgb = ORANGE
    bar.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.8), Inches(11.3), Inches(3.2))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "ACE — Intelligent Mobile Agent System"
    p0.font.size = Pt(36)
    p0.font.bold = True
    p0.font.color.rgb = WHITE

    p1 = tf1.add_paragraph()
    p1.text = "Simulasi Multi-Agent Customer Service Hotel dengan Migrasi State Mobile Agent Antar-Node Logis"
    p1.font.size = Pt(18)
    p1.font.color.rgb = RGBColor(203, 213, 225)
    p1.space_before = Pt(14)

    p2 = tf1.add_paragraph()
    p2.text = "Mata Kuliah: Agent Enterprise  |  Program Studi Informatika / Sistem Informasi"
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.color.rgb = ORANGE
    p2.space_before = Pt(20)

    # Team Card on Title Slide
    team_card = add_card(s1, Inches(1.0), Inches(5.2), Inches(11.3), Inches(1.5), bg_color=RGBColor(33, 46, 70), border_color=ORANGE)
    tb_team = s1.shapes.add_textbox(Inches(1.2), Inches(5.3), Inches(10.9), Inches(1.3))
    tf_team = tb_team.text_frame
    p_t0 = tf_team.paragraphs[0]
    p_t0.text = "Disusun oleh: KELOMPOK 5"
    p_t0.font.size = Pt(12)
    p_t0.font.bold = True
    p_t0.font.color.rgb = ORANGE

    p_t1 = tf_team.add_paragraph()
    p_t1.text = "• M. Al lail Qadrillah     • Dimas Prabowo     • Monanta Alfiareza     • Frans Alwan"
    p_t1.font.size = Pt(14)
    p_t1.font.color.rgb = WHITE
    p_t1.space_before = Pt(8)

    p_t2 = tf_team.add_paragraph()
    p_t2.text = "Hotel Fiktif: Hotel Nusantara Demo  |  100% Offline, Deterministic, & Local Simulation"
    p_t2.font.size = Pt(11)
    p_t2.font.color.rgb = RGBColor(148, 163, 184)
    p_t2.space_before = Pt(6)

    # ==========================================
    # SLIDE 2: Latar Belakang & Masalah
    # ==========================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2, LIGHT_BG)
    add_header(s2, "Latar Belakang: Tantangan Operasional Customer Service Hotel")

    # 3 Cards Problem
    problems = [
        ("Tantangan Integrasi Data", "Data Front Office (reservasi & tamu) terpisah secara logis dan fisik dari data Operations (housekeeping & kesiapan teknis kamar). Integrasi manual sering memicu delay dan salah kamar."),
        ("Keterbatasan Pendekatan Statis", "Pada arsitektur agen tradisional, agen pusat harus meminta seluruh data database untuk diinspeksi ke node pusat, menimbulkan pemborosan bandwidth dan risiko kebocoran data sensitif."),
        ("Solusi: Konsep Mobile Agent", "Memungkinkan agen cerdas (Mobile Investigator) berpindah lokasi eksekusi langsung ke node tempat data berada, menginspeksi secara lokal, dan hanya mengirimkan hasil akhir.")
    ]

    for idx, (title, desc) in enumerate(problems):
        x = Inches(0.8 + idx * 4.0)
        add_card(s2, x, Inches(1.8), Inches(3.7), Inches(4.8))
        tb = s2.shapes.add_textbox(x + Inches(0.2), Inches(2.0), Inches(3.3), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = f"0{idx+1}"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = ORANGE

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(16)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY
        p_title.space_before = Pt(10)

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(13)
        p_desc.font.color.rgb = DARK_TEXT
        p_desc.space_before = Pt(14)

    # ==========================================
    # SLIDE 3: Paradigma Mobile vs Static Agent
    # ==========================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3, LIGHT_BG)
    add_header(s3, "Paradigma: Mobile Agent vs Static Agent")

    # 2 Comparison Columns
    # Left: Static
    add_card(s3, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_st = s3.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.4))
    tf_st = tb_st.text_frame
    tf_st.word_wrap = True
    p = tf_st.paragraphs[0]
    p.text = "🏢 Static Agent (Baseline)"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = NAVY

    items_st = [
        "Agen investigator menetap di Front Office.",
        "Mengirim batch RPC request 'INSPECT_CANDIDATES' ke Operations Agent.",
        "Data ditarik melintasi jaringan/node.",
        "0 Migrasi agen terjadi.",
        "Hasil bisnis sama, tetapi bergantung pada transfer data bolak-balik."
    ]
    for item in items_st:
        p_i = tf_st.add_paragraph()
        p_i.text = "• " + item
        p_i.font.size = Pt(13)
        p_i.font.color.rgb = DARK_TEXT
        p_i.space_before = Pt(10)

    # Right: Mobile
    add_card(s3, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8), border_color=ORANGE)
    tb_mb = s3.shapes.add_textbox(Inches(7.2), Inches(2.0), Inches(5.0), Inches(4.4))
    tf_mb = tb_mb.text_frame
    tf_mb.word_wrap = True
    p = tf_mb.paragraphs[0]
    p.text = "🚀 Mobile Agent (Pendekatan Proyek)"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = ORANGE

    items_mb = [
        "Agen Mobile Investigator mengepak state ke Checkpoint JSON kanonik.",
        "Instance lama ditutup (Depart), instance baru lahir di node Operations (Arrive).",
        "Inspeksi dilakukan secara lokal di node data Operations.",
        "1 Migrasi sukses dengan kontinuitas ID ('MA-001').",
        "Hanya hasil akhir yang dilaporkan kembali ke Front Office."
    ]
    for item in items_mb:
        p_i = tf_mb.add_paragraph()
        p_i.text = "• " + item
        p_i.font.size = Pt(13)
        p_i.font.color.rgb = DARK_TEXT
        p_i.space_before = Pt(10)

    # ==========================================
    # SLIDE 4: Arsitektur 2 Node Logis & 7 Agen
    # ==========================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4, LIGHT_BG)
    add_header(s4, "Arsitektur Sistem: Dua Node Logis & Spesialisasi Agen")

    # Node 1: Front Office
    add_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_fo = s4.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(5.0), Inches(4.4))
    tf_fo = tb_fo.text_frame
    tf_fo.word_wrap = True
    p = tf_fo.paragraphs[0]
    p.text = "🏢 Node FRONT_OFFICE"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = ACCENT_BLUE

    p_sub = tf_fo.add_paragraph()
    p_sub.text = "Database In-Memory: Data Tamu, Reservasi, Folio, FAQ"
    p_sub.font.size = Pt(11)
    p_sub.font.color.rgb = MUTED_TEXT
    p_sub.space_before = Pt(4)

    agents_fo = [
        ("Scenario Scout", "Menerima pesan tamu & menyusun request terstruktur."),
        ("Orchestrator", "Pusat koordinasi, triase policy/ML, & manajemen kasus."),
        ("Reservation Agent", "Mencari kandidat kamar & eksekusi commit mutasi kamar."),
        ("Billing Agent", "Membaca folio tagihan & kalkulasi biaya minibar/kamar."),
        ("Concierge Agent", "Menjawab FAQ hotel (check-in, check-out, Wi-Fi).")
    ]
    for name, desc in agents_fo:
        p_a = tf_fo.add_paragraph()
        p_a.text = f"• {name}: {desc}"
        p_a.font.size = Pt(12)
        p_a.font.color.rgb = DARK_TEXT
        p_a.space_before = Pt(8)

    # Node 2: Operations
    add_card(s4, Inches(6.9), Inches(1.8), Inches(5.6), Inches(4.8))
    tb_op = s4.shapes.add_textbox(Inches(7.2), Inches(2.0), Inches(5.0), Inches(4.4))
    tf_op = tb_op.text_frame
    tf_op.word_wrap = True
    p = tf_op.paragraphs[0]
    p.text = "🔧 Node OPERATIONS"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = GREEN

    p_sub2 = tf_op.add_paragraph()
    p_sub2.text = "Database In-Memory: Status Kesiapan Kamar & Tiket Petugas"
    p_sub2.font.size = Pt(11)
    p_sub2.font.color.rgb = MUTED_TEXT
    p_sub2.space_before = Pt(4)

    agents_op = [
        ("Operations Agent", "Mengelola kesiapan fisik kamar, pemblokiran maintenance, & tiket staf."),
        ("Mobile Investigator (MA-001)", "Agen yang bermigrasi dari Front Office ke Operations untuk menginspeksi kesiapan kamar langsung di sumber data.")
    ]
    for name, desc in agents_op:
        p_a = tf_op.add_paragraph()
        p_a.text = f"• {name}: {desc}"
        p_a.font.size = Pt(12)
        p_a.font.color.rgb = DARK_TEXT
        p_a.space_before = Pt(10)

    p_note = tf_op.add_paragraph()
    p_note.text = "🔒 Kontrol Akses: Agen tidak dapat mengakses data di luar kapabilitas node tempat ia aktif saat ini."
    p_note.font.size = Pt(11)
    p_note.font.bold = True
    p_note.font.color.rgb = ORANGE
    p_note.space_before = Pt(24)

    # ==========================================
    # SLIDE 5: Protokol Migrasi & Integritas
    # ==========================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5, LIGHT_BG)
    add_header(s5, "Protokol Migrasi State Agen & Jaminan Keamanan")

    stages = [
        ("1. PREPARE", "Serialisasi state agen (goal, kandidat kamar, constraints) menjadi Checkpoint JSON kanonik UTF-8. Menghasilkan hash SHA-256."),
        ("2. DEPART", "Instance agen di node asal ditutup dan diarsipkan. Instance aktif = 0; agen hanya berada di jalur transit."),
        ("3. ARRIVE", "Instance baru dibuat di node tujuan dari checkpoint. Integritas hash SHA-256 & skema divalidasi sebelum aktivasi."),
        ("4. ROLLBACK", "Jika terjadi anomali (hash rusak, skema tidak cocok), mekanisme rollback memulihkan agen ke node asal secara aman.")
    ]

    for idx, (title, desc) in enumerate(stages):
        x = Inches(0.8 + idx * 3.0)
        add_card(s5, x, Inches(1.8), Inches(2.7), Inches(4.8))
        tb = s5.shapes.add_textbox(x + Inches(0.15), Inches(2.0), Inches(2.4), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = ORANGE

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = DARK_TEXT
        p_desc.space_before = Pt(12)

    # ==========================================
    # SLIDE 6: Pengalaman Multi-Stakeholder
    # ==========================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6, LIGHT_BG)
    add_header(s6, "Antarmuka Berbasis Peran (Multi-Stakeholder Experience)")

    stakeholders = [
        ("🛎️ Tamu Hotel (Customer)", "Portal interaktif ramah pengguna:\n• Chat asisten AI hotel\n• Kartu penawaran kamar baru (bersih & setara)\n• Tombol persetujuan tamu (Consent) langsung\n• Pilihan layanan cepat & pesan bebas"),
        ("👔 Staf Hotel (Operations)", "Dashboard operasional terpadu:\n• Ringkasan KPI tiket tugas & eskalasi\n• Inbox eskalasi kasus yang butuh staf\n• Tombol 'Ambil Alih' & resolusi catatan staf\n• Papan kerja tiket housekeeping/maintenance"),
        ("💻 Tim IT / DevOps", "Konsol observabilitas mendalam:\n• Pemantauan topologi agen & node\n• Verifikasi integritas hash SHA-256\n• Pengukuran payload bytes & latency\n• Unduh laporan audit log JSON lengkap")
    ]

    for idx, (title, desc) in enumerate(stakeholders):
        x = Inches(0.8 + idx * 4.0)
        add_card(s6, x, Inches(1.8), Inches(3.7), Inches(4.8))
        tb = s6.shapes.add_textbox(x + Inches(0.2), Inches(2.0), Inches(3.3), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.color.rgb = NAVY

        p_desc = tf.add_paragraph()
        p_desc.text = desc
        p_desc.font.size = Pt(12)
        p_desc.font.color.rgb = DARK_TEXT
        p_desc.space_before = Pt(14)

    # ==========================================
    # SLIDE 7: Skenario Pengujian (S01 - S09)
    # ==========================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7, LIGHT_BG)
    add_header(s7, "9 Skenario Pengujian Realistis & Otomatisasi 1-Klik")

    scenarios_summary = [
        ("S01: AC Rusak — Kamar Setara", "Migrasi Mobile Investigator -> Temukan R103 ready -> Proposal -> Persetujuan Tamu -> Selesai Digital."),
        ("S02: AC Rusak — Minta Upgrade", "Kamar setara habis. Bukti terkumpul, namun dialihkan ke staf karena aturan wajib eskalasi upgrade."),
        ("S03: Sengketa Tagihan Minibar", "Tamu menolak tagihan F02. Kebijakan finansial wajib eskalasi staf (WAITING_HUMAN)."),
        ("S04: Tanya Jam Check-In", "Pertanyaan jam operasional. Dijawab instan oleh Concierge Agent dari FAQ lokal (14.00)."),
        ("S05: Permintaan Handuk", "Layanan kamar. Tiket Housekeeping dibuat otomatis (PENDING -> IN_PROGRESS -> DONE)."),
        ("S06: Rincian Tagihan", "Menampilkan folio tagihan terperinci (Total Rp650.000) tanpa perubahan saldo."),
        ("S07: Jam Check-Out & Late", "Informasi check-out maksimal 12.00 WIB dan alur permohonan late check-out."),
        ("S08: Permintaan Bantal Tambahan", "Housekeeping membuat work order antar bantal tambahan ke kamar tamu."),
        ("S09: Informasi Wi-Fi Hotel", "Concierge menjawab detail SSID HotelNusantara_Guest tanpa kata sandi.")
    ]

    for idx, (title, desc) in enumerate(scenarios_summary):
        col = idx % 3
        row = idx // 3
        x = Inches(0.8 + col * 4.0)
        y = Inches(1.8 + row * 1.6)

        add_card(s7, x, y, Inches(3.7), Inches(1.4))
        tb = s7.shapes.add_textbox(x + Inches(0.15), y + Inches(0.1), Inches(3.4), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = ORANGE

        p_d = tf.add_paragraph()
        p_d.text = desc
        p_d.font.size = Pt(10)
        p_d.font.color.rgb = DARK_TEXT
        p_d.space_before = Pt(3)

    # ==========================================
    # SLIDE 8: Hasil Evaluasi & Kesetaraan
    # ==========================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8, LIGHT_BG)
    add_header(s8, "Evaluasi & Benchmark: Kesetaraan Hasil Bisnis")

    add_card(s8, Inches(0.8), Inches(1.8), Inches(11.7), Inches(4.8))
    tb_ev = s8.shapes.add_textbox(Inches(1.1), Inches(2.0), Inches(11.1), Inches(4.4))
    tf_ev = tb_ev.text_frame
    tf_ev.word_wrap = True

    p = tf_ev.paragraphs[0]
    p.text = "📊 Temuan Utama Evaluasi Komparatif (S01 & S02)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = NAVY

    eval_points = [
        ("Hasil Bisnis 100% Identik", "Kedua jalur (Mobile dan Static) menghasilkan keputusan kamar, tarif akhir, dan status tiket yang sama persis karena berbagi kernel evaluasi murni ('evaluate_candidates')."),
        ("Efisiensi Payload & Trafik Komunikasi", "Jalur Mobile Agent mengemas migrasi dalam satu checkpoint kanonik. Tidak ada polling berulang atau penarikan seluruh database reservasi/kamar."),
        ("Keandalan & Kebijakan (Policy vs ML)", "Prediksi Machine Learning (Logistic Regression) dikombinasikan dengan Mandatory Rules. Jika terdeteksi risiko tinggi/sengketa/permintaan staf, aturan bisnis selalu mengesampingkan ML."),
        ("Pengujian Otomatis Menyeluruh", "Seluruh 68 automated unit & integration tests lulus (100% pass) dalam waktu < 4 detik secara deterministik dan offline.")
    ]

    for title, desc in eval_points:
        p_t = tf_ev.add_paragraph()
        p_t.text = "✓ " + title + ": "
        p_t.font.size = Pt(13)
        p_t.font.bold = True
        p_t.font.color.rgb = ORANGE
        p_t.space_before = Pt(12)

        # append description
        run = p_t.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = DARK_TEXT

    # ==========================================
    # SLIDE 9: Kesimpulan & Penutup
    # ==========================================
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9, NAVY)

    tb9 = s9.shapes.add_textbox(Inches(1.5), Inches(1.8), Inches(10.3), Inches(4.0))
    tf9 = tb9.text_frame
    tf9.word_wrap = True

    p0 = tf9.paragraphs[0]
    p0.text = "Kesimpulan & Nilai Tambah Enterprise"
    p0.font.size = Pt(32)
    p0.font.bold = True
    p0.font.color.rgb = WHITE

    points = [
        "Membuktikan kelayakan implementasi konsep Mobile Agent pada domain perhotelan modern.",
        "Menjaga desentralisasi data tanpa mengorbankan kecepatan penyelesaian keluhan tamu.",
        "Menyediakan antarmuka intuitif untuk 3 stakeholder utama: Tamu, Staf Hotel, dan Tim IT.",
        "Sistem siap didemonstrasikan secara langsung (live interactive demo) maupun otomatis (1-click)."
    ]

    for pt in points:
        p = tf9.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(16)
        p.font.color.rgb = RGBColor(226, 232, 240)
        p.space_before = Pt(14)

    p_close = tf9.add_paragraph()
    p_close.text = "Terima Kasih — Kelompok 5 (Agent Enterprise)"
    p_close.font.size = Pt(20)
    p_close.font.bold = True
    p_close.font.color.rgb = ORANGE
    p_close.space_before = Pt(28)

    # Save presentation
    output_path = "presentasi_ace_mobile_agent.pptx"
    prs.save(output_path)
    print(f"Presentation saved to {os.path.abspath(output_path)}")

if __name__ == "__main__":
    create_deck()
